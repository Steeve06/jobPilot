import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


def render_resume_to_docx(data):
    doc = Document()
    _set_page_margins(doc)
    _set_base_styles(doc)

    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _tighten(name_para, space_after=2)
    name_run = name_para.add_run(data.get('full_name') or 'Your Name')
    name_run.bold = True
    name_run.font.size = Pt(17)

    contact_parts = [
        data.get('email'), data.get('phone'), data.get('location'),
        data.get('linkedin_url'), data.get('github_url'),
    ]
    contact_line = ' | '.join(p for p in contact_parts if p)
    if contact_line:
        contact_para = doc.add_paragraph(contact_line)
        contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _tighten(contact_para, space_after=8)
        contact_para.runs[0].font.size = Pt(9)

    if data.get('summary'):
        _add_section_heading(doc, 'Summary')
        p = doc.add_paragraph(data['summary'])
        _tighten(p, space_after=6)

    if data.get('skills'):
        _add_section_heading(doc, 'Skills')
        p = doc.add_paragraph(', '.join(data['skills']))
        _tighten(p, space_after=6)

    if data.get('experiences'):
        _add_section_heading(doc, 'Experience')
        for exp in data['experiences']:
            _add_experience_entry(doc, exp)

    if data.get('projects'):
        _add_section_heading(doc, 'Projects')
        for proj in data['projects']:
            _add_project_entry(doc, proj)

    if data.get('education'):
        _add_section_heading(doc, 'Education')
        for edu in data['education']:
            line = f"{edu.get('degree', '')} — {edu.get('school', '')}".strip(' —')
            p = doc.add_paragraph(line)
            _tighten(p, space_after=2)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _set_page_margins(doc):
    section = doc.sections[0]
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)


def _set_base_styles(doc):
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.line_spacing = 1.0


def _tighten(paragraph, space_before=0, space_after=4):
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.space_after = Pt(space_after)


def _add_section_heading(doc, text):
    heading = doc.add_heading(text, level=2)
    heading.style.font.size = Pt(12)
    _tighten(heading, space_before=6, space_after=2)


def _add_experience_entry(doc, exp):
    header_para = doc.add_paragraph()
    _tighten(header_para, space_after=0)
    title_run = header_para.add_run(exp.get('title', ''))
    title_run.bold = True

    date_range = f"{exp.get('start_date') or ''} – {exp.get('end_date') or 'Present'}"
    header_para.add_run(f'\t{date_range}')
    header_para.paragraph_format.tab_stops.add_tab_stop(
        doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin,
        alignment=WD_ALIGN_PARAGRAPH.RIGHT,
    )

    company_para = doc.add_paragraph(exp.get('company', ''))
    _tighten(company_para, space_after=3)
    company_para.runs[0].italic = True

    for bullet in exp.get('bullets', []):
        p = doc.add_paragraph(bullet.get('text', ''), style='List Bullet')
        _tighten(p, space_after=1)


def _add_project_entry(doc, proj):
    name_para = doc.add_paragraph()
    _tighten(name_para, space_after=0)
    name_para.add_run(proj.get('name', '')).bold = True
    for bullet in proj.get('bullets', []):
        p = doc.add_paragraph(bullet.get('text', ''), style='List Bullet')
        _tighten(p, space_after=1)